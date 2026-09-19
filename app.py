import streamlit as st
from quiz_chain import answer_sports_question, get_sports_news, generate_quiz
from fpdf import FPDF
from docx import Document
import io
import time
import re

st.set_page_config(
    page_title="SportsIQ AI",
    page_icon="🏏",
    layout="wide",
    initial_sidebar_state="collapsed"
)

# Session State Init
if "active_tab" not in st.session_state:
    st.session_state["active_tab"] = 0
if "quiz_history" not in st.session_state:
    st.session_state["quiz_history"] = []
if "questions" not in st.session_state:
    st.session_state["questions"] = []
if "user_answers" not in st.session_state:
    st.session_state["user_answers"] = {}
if "quiz_submitted" not in st.session_state:
    st.session_state["quiz_submitted"] = False
if "quiz_mode" not in st.session_state:
    st.session_state["quiz_mode"] = "Attempt Online"
if "start_time" not in st.session_state:
    st.session_state["start_time"] = None
if "quiz_sport" not in st.session_state:
    st.session_state["quiz_sport"] = ""

st.markdown("""
<style>
@import url('https://fonts.googleapis.com/css2?family=Orbitron:wght@700&family=Exo+2:wght@400;600;700&display=swap');
* { font-family: 'Exo 2', sans-serif; }
header[data-testid="stHeader"] { display: none; }
#MainMenu { visibility: hidden; }
footer { visibility: hidden; }
.stApp {
    background-image: linear-gradient(rgba(15,23,42,0.93), rgba(15,23,42,0.96)),
    url('https://images.unsplash.com/photo-1540747913346-19e32dc3e97e?w=1920&q=80');
    background-size: cover;
    background-position: center;
    background-attachment: fixed;
    color: #F8FAFC;
}
.navbar {
    background-color: #1E293B;
    padding: 15px 40px;
    display: flex;
    justify-content: space-between;
    align-items: center;
    border-bottom: 2px solid #22C55E;
    margin-bottom: 10px;
    position: sticky;
    top: 0;
    z-index: 999;
}
.navbar-links { display: flex; gap: 25px; align-items: center; }
.main-title {
    text-align: center;
    font-size: 2.8em;
    font-weight: 700;
    color: #22C55E;
    font-family: 'Orbitron', sans-serif;
    padding: 20px 0 5px 0;
    text-shadow: 0 0 30px rgba(34,197,94,0.3);
}
.sub-title {
    text-align: center;
    color: #94A3B8;
    font-size: 1.1em;
    margin-bottom: 20px;
}
.section-header {
    font-size: 1.4em;
    font-weight: 700;
    color: #FACC15;
    font-family: 'Orbitron', sans-serif;
    padding: 15px 0 10px 0;
    border-bottom: 2px solid #22C55E;
    margin-bottom: 20px;
}
.answer-box {
    background-color: #1E293B;
    border-left: 4px solid #22C55E;
    padding: 20px;
    border-radius: 8px;
    color: #F8FAFC;
    margin: 10px 0;
    line-height: 1.7;
}
.news-card {
    background-color: #1E293B;
    border-left: 4px solid #FACC15;
    padding: 15px;
    border-radius: 8px;
    margin: 10px 0;
}
.news-title {
    color: #FACC15;
    font-size: 1.05em;
    font-weight: 700;
    margin-bottom: 8px;
}
.news-summary { color: #94A3B8; font-size: 0.9em; line-height: 1.6; }
.confidence-box {
    background-color: #1E293B;
    border: 1px solid #22C55E;
    padding: 10px 20px;
    border-radius: 8px;
    text-align: center;
    margin: 5px;
}
.related-box {
    background-color: #1E293B;
    border-left: 4px solid #94A3B8;
    padding: 15px;
    border-radius: 8px;
    margin: 10px 0;
    color: #94A3B8;
    line-height: 1.8;
}
.question-card {
    background-color: #1E293B;
    border: 1px solid #334155;
    padding: 20px;
    border-radius: 10px;
    margin: 15px 0;
}
.correct-ans {
    background-color: rgba(34,197,94,0.15);
    border-left: 4px solid #22C55E;
    padding: 10px 15px;
    border-radius: 5px;
    margin: 5px 0;
    color: #22C55E;
}
.wrong-ans {
    background-color: rgba(239,68,68,0.1);
    border-left: 4px solid #EF4444;
    padding: 10px 15px;
    border-radius: 5px;
    margin: 5px 0;
    color: #EF4444;
}
.neutral-ans { padding: 10px 15px; margin: 5px 0; color: #94A3B8; }
.score-box {
    background: linear-gradient(135deg, #1E293B, #0F172A);
    border: 2px solid #22C55E;
    padding: 30px;
    border-radius: 15px;
    text-align: center;
    margin: 20px 0;
}
.timer-box {
    background-color: #1E293B;
    border: 2px solid #FACC15;
    padding: 10px 20px;
    border-radius: 8px;
    text-align: center;
    font-size: 1.5em;
    font-weight: 700;
    color: #FACC15;
    margin-bottom: 15px;
}
.footer {
    background-color: #1E293B;
    padding: 40px 20px 20px 20px;
    margin-top: 60px;
    border-top: 2px solid #22C55E;
}
div.stButton > button {
    background-color: #22C55E;
    color: #0F172A;
    font-weight: 700;
    border-radius: 8px;
    border: none;
    padding: 8px 20px;
    width: 100%;
}
div.stButton > button:hover { background-color: #FACC15; color: #0F172A; }
label { color: #F8FAFC !important; }
p { color: #F8FAFC; }
h1, h2, h3 { color: #F8FAFC; }
.stTabs [data-baseweb="tab"] { color: #94A3B8; }
.stTabs [aria-selected="true"] { color: #22C55E !important; }
.nav-btn button {
    background: transparent !important;
    color: #F8FAFC !important;
    border: none !important;
    font-weight: 600 !important;
    font-size: 0.9em !important;
    padding: 5px 10px !important;
    width: auto !important;
}
.nav-btn button:hover { color: #22C55E !important; background: transparent !important; }
@media (max-width: 768px) {
    .navbar { flex-direction: column; gap: 10px; padding: 10px 20px; }
    .navbar-links { flex-wrap: wrap; gap: 15px; justify-content: center; }
    .main-title { font-size: 1.8em !important; }
    .footer > div { grid-template-columns: 1fr 1fr !important; }
}
@media (max-width: 480px) {
    .main-title { font-size: 1.3em !important; }
    .footer > div { grid-template-columns: 1fr !important; }
}
</style>
""", unsafe_allow_html=True)

# ─────────────────────────────────────────
# NAVBAR
# ─────────────────────────────────────────
st.markdown("""
<div class="navbar">
    <div class="navbar-brand">
        <span style="color:#22C55E; font-size:1.5em; font-weight:700; font-family:'Orbitron',sans-serif;">🏏 SportsIQ AI</span><br>
        <span style="color:#94A3B8; font-size:0.8em;">AI-Powered Sports Assistant</span><br>
        <span style="color:#FACC15; font-size:0.75em; font-weight:600;">Ask • Learn • Compete</span>
    </div>
</div>
""", unsafe_allow_html=True)

# Navbar navigation buttons
nav1, nav2, nav3, nav4, nav5 = st.columns(5)
with nav1:
    st.markdown('<div class="nav-btn">', unsafe_allow_html=True)
    if st.button("🏠 Home"):
        st.session_state["active_tab"] = 0
        st.rerun()
    st.markdown('</div>', unsafe_allow_html=True)
with nav2:
    st.markdown('<div class="nav-btn">', unsafe_allow_html=True)
    if st.button("💬 Ask AI"):
        st.session_state["active_tab"] = 0
        st.rerun()
    st.markdown('</div>', unsafe_allow_html=True)
with nav3:
    st.markdown('<div class="nav-btn">', unsafe_allow_html=True)
    if st.button("🎯 Quiz Arena"):
        st.session_state["active_tab"] = 1
        st.rerun()
    st.markdown('</div>', unsafe_allow_html=True)
with nav4:
    st.markdown('<div class="nav-btn">', unsafe_allow_html=True)
    if st.button("📰 Sports News"):
        st.session_state["active_tab"] = 2
        st.rerun()
    st.markdown('</div>', unsafe_allow_html=True)
with nav5:
    st.markdown('<div class="nav-btn">', unsafe_allow_html=True)
    if st.button("📚 Quiz History"):
        st.session_state["active_tab"] = 3
        st.rerun()
    st.markdown('</div>', unsafe_allow_html=True)

st.markdown("---")

# ─────────────────────────────────────────
# TITLE
# ─────────────────────────────────────────
st.markdown('<div class="main-title">🏏 SportsIQ AI</div>', unsafe_allow_html=True)
st.markdown('<div class="sub-title">Your Personal Sports Knowledge Hub | Powered by AI</div>', unsafe_allow_html=True)
st.markdown("---")

# ─────────────────────────────────────────
# HELPER FUNCTIONS
# ─────────────────────────────────────────
def clean_text(text):
    if not text:
        return ""
    return text.encode("latin-1", errors="replace").decode("latin-1")

def parse_quiz(raw):
    questions = []
    blocks = raw.strip().split("---")
    for block in blocks:
        block = block.strip()
        if not block:
            continue
        q = {
            "question": "", "difficulty": "Medium",
            "options": [], "answer": "",
            "why_correct": "", "why_a": "",
            "why_b": "", "why_c": "", "why_d": ""
        }
        for line in block.split("\n"):
            line = line.strip()
            if line.startswith("QUESTION:"):
                q["question"] = line.replace("QUESTION:", "").strip()
            elif line.startswith("DIFFICULTY:"):
                q["difficulty"] = line.replace("DIFFICULTY:", "").strip()
            elif line.startswith(("A)", "B)", "C)", "D)")):
                q["options"].append(line)
            elif line.startswith("ANSWER:"):
                q["answer"] = line.replace("ANSWER:", "").strip()
            elif line.startswith("WHY_CORRECT:"):
                q["why_correct"] = line.replace("WHY_CORRECT:", "").strip()
            elif line.startswith("WHY_A:"):
                q["why_a"] = line.replace("WHY_A:", "").strip()
            elif line.startswith("WHY_B:"):
                q["why_b"] = line.replace("WHY_B:", "").strip()
            elif line.startswith("WHY_C:"):
                q["why_c"] = line.replace("WHY_C:", "").strip()
            elif line.startswith("WHY_D:"):
                q["why_d"] = line.replace("WHY_D:", "").strip()
        if q["question"] and q["options"]:
            questions.append(q)
    return questions

def generate_pdf(questions, user_answers=None, show_answers=True):
    pdf = FPDF()
    pdf.set_margins(15, 15, 15)
    pdf.add_page()
    pdf.set_auto_page_break(auto=True, margin=15)

    def safe_write(h, text):
        try:
            pdf.multi_cell(0, h, clean_text(str(text)))
        except Exception:
            pdf.multi_cell(0, h, "[text unavailable]")

    pdf.set_font("Helvetica", "B", 14)
    safe_write(10, "SportsIQ AI - Sports Quiz")
    pdf.set_font("Helvetica", "", 9)
    safe_write(8, "Generated by SportsIQ AI | Kishore Reddy Dudipala")
    pdf.ln(4)

    for i, q in enumerate(questions):
        pdf.set_font("Helvetica", "B", 10)
        safe_write(7, f"Q{i+1} [{q.get('difficulty','Medium')}]: {q['question']}")
        pdf.set_font("Helvetica", "", 9)
        for opt in q['options']:
            safe_write(6, f"  {opt}")
        if show_answers and user_answers is not None:
            user_ans = user_answers.get(i, "Not Attempted")
            correct = q['answer']
            status = "CORRECT" if user_ans == correct else "WRONG"
            pdf.set_font("Helvetica", "B", 9)
            safe_write(6, f"Your Answer: {user_ans} | Correct: {correct} | {status}")
            pdf.set_font("Helvetica", "", 8)
            safe_write(5, f"Explanation: {q.get('why_correct', '')}")
            safe_write(5, f"Why A: {q.get('why_a', '')}")
            safe_write(5, f"Why B: {q.get('why_b', '')}")
            safe_write(5, f"Why C: {q.get('why_c', '')}")
            safe_write(5, f"Why D: {q.get('why_d', '')}")
        elif show_answers:
            pdf.set_font("Helvetica", "B", 9)
            safe_write(6, f"Answer: {q['answer']}")
            pdf.set_font("Helvetica", "", 8)
            safe_write(5, f"Explanation: {q.get('why_correct', '')}")
            safe_write(5, f"Why A: {q.get('why_a', '')}")
            safe_write(5, f"Why B: {q.get('why_b', '')}")
            safe_write(5, f"Why C: {q.get('why_c', '')}")
            safe_write(5, f"Why D: {q.get('why_d', '')}")
        pdf.ln(3)
    return bytes(pdf.output())

def generate_word(questions, user_answers=None, show_answers=True):
    doc = Document()
    doc.add_heading("SportsIQ AI - Sports Quiz", 0)
    doc.add_paragraph("Generated by SportsIQ AI | Kishore Reddy Dudipala")
    for i, q in enumerate(questions):
        doc.add_heading(f"Q{i+1} [{q.get('difficulty','Medium')}]: {q['question']}", level=2)
        for opt in q['options']:
            doc.add_paragraph(opt, style='List Bullet')
        if show_answers and user_answers is not None:
            user_ans = user_answers.get(i, "Not Attempted")
            correct = q['answer']
            status = "CORRECT" if user_ans == correct else "WRONG"
            doc.add_paragraph(f"Your Answer: {user_ans} | Correct: {correct} | {status}")
            doc.add_paragraph(f"Explanation: {q.get('why_correct', '')}")
            doc.add_paragraph(f"Why A: {q.get('why_a', '')}")
            doc.add_paragraph(f"Why B: {q.get('why_b', '')}")
            doc.add_paragraph(f"Why C: {q.get('why_c', '')}")
            doc.add_paragraph(f"Why D: {q.get('why_d', '')}")
        elif show_answers:
            doc.add_paragraph(f"Answer: {q['answer']}")
            doc.add_paragraph(f"Explanation: {q.get('why_correct', '')}")
            doc.add_paragraph(f"Why A: {q.get('why_a', '')}")
            doc.add_paragraph(f"Why B: {q.get('why_b', '')}")
            doc.add_paragraph(f"Why C: {q.get('why_c', '')}")
            doc.add_paragraph(f"Why D: {q.get('why_d', '')}")
        doc.add_paragraph("")
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.getvalue()

# ─────────────────────────────────────────
# TABS
# ─────────────────────────────────────────
tab1, tab2, tab3, tab4 = st.tabs(["💬 Sports Q&A", "🎯 Quiz Arena", "📰 Sports News", "📚 Quiz History"])

# ═══════════════════════════════════════════
# TAB 1 — Q&A
# ═══════════════════════════════════════════
with tab1:
    st.markdown('<div class="section-header">💬 Sports Q&A — Ask Anything</div>', unsafe_allow_html=True)
    question = st.text_input("", placeholder="e.g. Who won IPL 2024? Who is the best cricketer in India?")

    if st.button("Ask AI 🔍", key="ask_btn"):
        if question.strip() == "":
            st.warning("Please enter a question!")
        else:
            with st.spinner("Searching web and analysing..."):
                response, search_result, sources = answer_sports_question(question)

            answer = ""
            related = []
            confidence = 85

            if "ANSWER:" in response:
                parts = response.split("ANSWER:")
                rest = parts[1] if len(parts) > 1 else ""
                if "RELATED QUESTIONS:" in rest:
                    ans_part, rel_part = rest.split("RELATED QUESTIONS:")
                    answer = ans_part.strip()
                    if "CONFIDENCE:" in rel_part:
                        rel_part, conf_part = rel_part.split("CONFIDENCE:")
                        try:
                            confidence = int(re.search(r'\d+', conf_part).group())
                        except:
                            confidence = 85
                    related = [r.strip().lstrip("123. ") for r in rel_part.strip().split("\n") if r.strip()]
                else:
                    answer = rest.strip()
            else:
                answer = response

            st.markdown(f'<div class="answer-box">{answer}</div>', unsafe_allow_html=True)

            col1, col2 = st.columns(2)
            with col1:
                color = "#22C55E" if confidence >= 80 else "#FACC15" if confidence >= 60 else "#EF4444"
                st.markdown(f'<div class="confidence-box">🎯 Answer Confidence: <span style="color:{color}; font-weight:700;">{confidence}%</span></div>', unsafe_allow_html=True)
            with col2:
                st.markdown(f'<div class="confidence-box">🔎 Search Sources Found: <span style="color:#22C55E; font-weight:700;">{sources}</span></div>', unsafe_allow_html=True)

            if related:
                related_html = "<br>".join([f"• {r}" for r in related if r])
                st.markdown(f'<div class="related-box"><strong style="color:#F8FAFC;">💡 You may also like:</strong><br><br>{related_html}</div>', unsafe_allow_html=True)

            with st.expander("🔎 Web Search Result Used"):
                st.markdown(f'<div style="color:#94A3B8; font-size:0.85em;">{search_result}</div>', unsafe_allow_html=True)

# ═══════════════════════════════════════════
# TAB 2 — QUIZ ARENA
# ═══════════════════════════════════════════
with tab2:
    st.markdown('<div class="section-header">🎯 Quiz Arena</div>', unsafe_allow_html=True)

    col1, col2, col3, col4 = st.columns(4)
    with col1:
        sport = st.selectbox("🏅 Sport", [
            "Cricket", "Football", "Badminton", "Hockey",
            "Kabaddi", "Tennis", "Basketball", "Wrestling",
            "Athletics", "Boxing", "Chess", "Kho Kho",
            "Volleyball", "Shooting", "Swimming",
            "All Indian Sports"
        ])
    with col2:
        num_questions = st.slider("📝 Questions", 5, 50, 10)
    with col3:
        time_limit = st.selectbox("⏱️ Time Limit", [
            "No Limit", "5 Minutes", "10 Minutes",
            "15 Minutes", "20 Minutes", "30 Minutes"
        ])
    with col4:
        mode = st.radio("📋 Mode", ["Attempt Online", "Download Blank"])

    if st.button("Generate Quiz 🎯", key="gen_quiz"):
        with st.spinner(f"Generating {num_questions} questions about {sport}... Please wait!"):
            raw = generate_quiz(sport, num_questions)
            questions = parse_quiz(raw)
            st.session_state["questions"] = questions
            st.session_state["user_answers"] = {}
            st.session_state["quiz_submitted"] = False
            st.session_state["quiz_mode"] = mode
            st.session_state["start_time"] = time.time()
            st.session_state["quiz_sport"] = sport
        st.success(f"✅ {len(questions)} questions generated!")

    # Download Blank
    if st.session_state["questions"] and st.session_state.get("quiz_mode") == "Download Blank":
        questions = st.session_state["questions"]
        st.markdown("### 📄 Download Options")
        col1, col2 = st.columns(2)
        with col1:
            st.markdown("**Blank Quiz**")
            st.download_button("📥 Blank PDF",
                data=generate_pdf(questions, show_answers=False),
                file_name="blank_quiz.pdf", mime="application/pdf", key="blank_pdf")
            st.download_button("📥 Blank Word",
                data=generate_word(questions, show_answers=False),
                file_name="blank_quiz.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                key="blank_word")
        with col2:
            st.markdown("**Quiz with Answers**")
            st.download_button("📥 Answers PDF",
                data=generate_pdf(questions, show_answers=True),
                file_name="quiz_answers.pdf", mime="application/pdf", key="ans_pdf")
            st.download_button("📥 Answers Word",
                data=generate_word(questions, show_answers=True),
                file_name="quiz_answers.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                key="ans_word")

    # Attempt Online
    if (st.session_state["questions"] and
        st.session_state.get("quiz_mode") == "Attempt Online" and
        not st.session_state.get("quiz_submitted", False)):

        questions = st.session_state["questions"]

        if time_limit != "No Limit":
            elapsed = time.time() - st.session_state.get("start_time", time.time())
            limit_secs = int(time_limit.split()[0]) * 60
            remaining = max(0, limit_secs - int(elapsed))
            mins, secs = divmod(remaining, 60)
            st.markdown(f'<div class="timer-box">⏱️ Time Remaining: {mins:02d}:{secs:02d}</div>', unsafe_allow_html=True)
            if remaining == 0:
                st.warning("⏰ Time is up!")
                st.session_state["quiz_submitted"] = True
                st.rerun()

        st.markdown(f"### 📋 {len(questions)} Questions | {st.session_state['quiz_sport']} Quiz")

        for i, q in enumerate(questions):
            diff_color = {"Easy": "#22C55E", "Medium": "#FACC15", "Hard": "#EF4444"}.get(q['difficulty'], "#94A3B8")
            st.markdown(f"""
            <div class="question-card">
                <strong style="color:#F8FAFC;">Q{i+1}: {q['question']}</strong>
                <span style="color:{diff_color}; font-size:0.8em; margin-left:10px;">[{q['difficulty']}]</span>
            </div>""", unsafe_allow_html=True)
            options_list = ["Not Attempted"] + q['options']
            selected = st.radio("", options_list, key=f"q_{i}", index=0)
            if selected != "Not Attempted":
                st.session_state["user_answers"][i] = selected[0]

        if st.button("Submit Quiz ✅", key="submit_quiz"):
            st.session_state["quiz_submitted"] = True
            st.rerun()

    # Results
    if st.session_state.get("quiz_submitted", False) and st.session_state["questions"]:
        questions = st.session_state["questions"]
        user_answers = st.session_state["user_answers"]
        score = sum(1 for i, q in enumerate(questions) if user_answers.get(i) == q['answer'])
        percentage = (score / len(questions) * 100) if questions else 0
        emoji = "🏆" if percentage >= 80 else "💪" if percentage >= 50 else "📚"

        st.markdown(f"""
        <div class="score-box">
            <h1 style="color:#22C55E;">{emoji} {score}/{len(questions)}</h1>
            <h3 style="color:#FACC15;">{percentage:.1f}% Score</h3>
            <p style="color:#94A3B8;">{"Excellent! You are a Sports Expert!" if percentage >= 80 else "Good effort! Keep learning!" if percentage >= 50 else "Keep practicing! You will get better!"}</p>
        </div>""", unsafe_allow_html=True)

        # Save to history
        st.session_state["quiz_history"].append({
            "sport": st.session_state.get("quiz_sport", "Unknown"),
            "score": score,
            "total": len(questions),
            "percentage": round(percentage, 1),
            "questions": questions,
            "user_answers": user_answers
        })

        st.markdown("### 📊 Detailed Results")
        for i, q in enumerate(questions):
            user_ans = user_answers.get(i, "Not Attempted")
            correct = q['answer']
            is_correct = user_ans == correct
            with st.expander(f"Q{i+1}: {q['question'][:60]}... | {'✅ Correct' if is_correct else '❌ Wrong'}"):
                for opt in q['options']:
                    opt_letter = opt[0]
                    if opt_letter == correct:
                        st.markdown(f'<div class="correct-ans">✅ {opt} ← Correct Answer</div>', unsafe_allow_html=True)
                    elif opt_letter == user_ans:
                        st.markdown(f'<div class="wrong-ans">❌ {opt} ← Your Answer</div>', unsafe_allow_html=True)
                    else:
                        st.markdown(f'<div class="neutral-ans">{opt}</div>', unsafe_allow_html=True)
                st.markdown(f"**✅ Why {correct} is correct:** {q.get('why_correct', '')}")
                st.markdown(f"**Why A:** {q.get('why_a', '')}")
                st.markdown(f"**Why B:** {q.get('why_b', '')}")
                st.markdown(f"**Why C:** {q.get('why_c', '')}")
                st.markdown(f"**Why D:** {q.get('why_d', '')}")

        st.markdown("### 📥 Download Results")
        col1, col2 = st.columns(2)
        with col1:
            st.download_button("📥 Results PDF",
                data=generate_pdf(questions, user_answers, show_answers=True),
                file_name="quiz_results.pdf", mime="application/pdf", key="res_pdf")
        with col2:
            st.download_button("📥 Results Word",
                data=generate_word(questions, user_answers, show_answers=True),
                file_name="quiz_results.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                key="res_word")

        if st.button("🔄 Try New Quiz", key="new_quiz"):
            st.session_state["questions"] = []
            st.session_state["user_answers"] = {}
            st.session_state["quiz_submitted"] = False
            st.session_state["start_time"] = None
            st.rerun()

# ═══════════════════════════════════════════
# TAB 3 — NEWS
# ═══════════════════════════════════════════
with tab3:
    st.markdown('<div class="section-header">📰 Latest Indian Sports News</div>', unsafe_allow_html=True)
    news_sport = st.selectbox("Select Sport", [
        "All Indian Sports", "Cricket", "Football",
        "Badminton", "Hockey", "Kabaddi",
        "Tennis", "Wrestling", "Athletics"
    ], key="news_sport")

    if st.button("Get Latest News 📰", key="get_news"):
        with st.spinner("Fetching latest sports news..."):
            news_raw = get_sports_news(news_sport)

        news_items = []
        current = {}
        for line in news_raw.split("\n"):
            line = line.strip()
            for n in ["1", "2", "3", "4"]:
                if line.startswith(f"NEWS{n}_TITLE:"):
                    if current:
                        news_items.append(current)
                    current = {"title": line.split(":", 1)[-1].strip(), "summary": "", "category": ""}
                elif line.startswith(f"NEWS{n}_SUMMARY:") and current:
                    current["summary"] = line.split(":", 1)[-1].strip()
                elif line.startswith(f"NEWS{n}_CATEGORY:") and current:
                    current["category"] = line.split(":", 1)[-1].strip()
        if current:
            news_items.append(current)

        if news_items:
            for item in news_items:
                st.markdown(f"""
                <div class="news-card">
                    <div class="news-title">📰 {item['title']}</div>
                    <div style="color:#22C55E; font-size:0.8em; margin-bottom:8px;">🏅 {item['category']}</div>
                    <div class="news-summary">{item['summary']}</div>
                </div>""", unsafe_allow_html=True)
        else:
            st.markdown(f'<div class="answer-box">{news_raw}</div>', unsafe_allow_html=True)

# ═══════════════════════════════════════════
# TAB 4 — QUIZ HISTORY
# ═══════════════════════════════════════════
with tab4:
    st.markdown('<div class="section-header">📚 Quiz History</div>', unsafe_allow_html=True)

    if not st.session_state["quiz_history"]:
        st.markdown('<div class="answer-box">No quizzes attempted yet. Go to Quiz Arena to start!</div>', unsafe_allow_html=True)
    else:
        st.markdown(f"**Total Quizzes Attempted: {len(st.session_state['quiz_history'])}**")
        for idx, entry in enumerate(reversed(st.session_state["quiz_history"])):
            emoji = "🏆" if entry['percentage'] >= 80 else "💪" if entry['percentage'] >= 50 else "📚"
            real_idx = len(st.session_state["quiz_history"]) - idx
            with st.expander(f"Quiz {real_idx}: {entry['sport']} | {entry['score']}/{entry['total']} ({entry['percentage']}%) {emoji}"):
                st.markdown(f"**Sport:** {entry['sport']}")
                st.markdown(f"**Score:** {entry['score']}/{entry['total']} ({entry['percentage']}%)")
                col1, col2 = st.columns(2)
                with col1:
                    st.download_button("📥 Download PDF",
                        data=generate_pdf(entry['questions'], entry['user_answers'], show_answers=True),
                        file_name=f"quiz_{entry['sport']}_{real_idx}.pdf",
                        mime="application/pdf",
                        key=f"hist_pdf_{idx}")
                with col2:
                    st.download_button("📥 Download Word",
                        data=generate_word(entry['questions'], entry['user_answers'], show_answers=True),
                        file_name=f"quiz_{entry['sport']}_{real_idx}.docx",
                        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                        key=f"hist_word_{idx}")

        if st.button("🗑️ Clear All History", key="clear_history"):
            st.session_state["quiz_history"] = []
            st.rerun()

# ─────────────────────────────────────────
# CTA
# ─────────────────────────────────────────
st.markdown("---")
st.markdown("""
<div style="text-align:center; padding:20px;">
    <h3 style="color:#F8FAFC;">Ready to Test Your Sports Knowledge?</h3>
    <p style="color:#94A3B8;">Sports Intelligence Powered by AI</p>
</div>
""", unsafe_allow_html=True)

# ─────────────────────────────────────────
# FOOTER
# ─────────────────────────────────────────
st.markdown("""
<div class="footer">
    <div style="display:grid; grid-template-columns:1fr 1fr 1fr 1fr; gap:20px; max-width:1100px; margin:0 auto;">
        <div>
            <h4 style="color:#22C55E; font-family:'Orbitron',sans-serif;">🏏 SportsIQ AI</h4>
            <p style="color:#94A3B8; font-size:0.85em;">AI-Powered Sports Assistant</p>
            <p style="color:#94A3B8; font-size:0.85em;">Ask • Learn • Compete</p>
            <p style="color:#FACC15; font-size:0.8em; margin-top:10px;">Sports Intelligence Powered by AI</p>
        </div>
        <div>
            <h4 style="color:#FACC15;">Quick Links</h4>
            <p style="color:#94A3B8; font-size:0.85em;">💬 Ask AI</p>
            <p style="color:#94A3B8; font-size:0.85em;">🎯 Quiz Arena</p>
            <p style="color:#94A3B8; font-size:0.85em;">📰 Sports News</p>
            <p style="color:#94A3B8; font-size:0.85em;">📚 Quiz History</p>
        </div>
        <div>
            <h4 style="color:#FACC15;">Developer</h4>
            <p style="color:#94A3B8; font-size:0.85em;">Kishore Reddy Dudipala</p>
            <p style="color:#94A3B8; font-size:0.85em;">B.E. Electronics & Communication Engineering</p>
            <p style="color:#94A3B8; font-size:0.85em;">Vasavi College of Engineering, Hyderabad</p>
        </div>
        <div>
            <h4 style="color:#FACC15;">Connect</h4>
            <p style="font-size:0.85em;">
                <a href="mailto:dkreddy2025@gmail.com" style="color:#22C55E; display:flex; align-items:center; gap:8px; margin-bottom:10px; text-decoration:none;">
                    <img src="https://cdn.jsdelivr.net/npm/simple-icons@v9/icons/gmail.svg" width="16" style="filter:invert(1) sepia(1) saturate(5) hue-rotate(90deg);"/>
                    dkreddy2025@gmail.com
                </a>
            </p>
            <p style="font-size:0.85em;">
                <a href="https://linkedin.com/in/kishorereddy945" style="color:#22C55E; display:flex; align-items:center; gap:8px; margin-bottom:10px; text-decoration:none;" target="_blank">
                    <img src="https://cdn.jsdelivr.net/npm/simple-icons@v9/icons/linkedin.svg" width="16" style="filter:invert(1) sepia(1) saturate(5) hue-rotate(90deg);"/>
                    linkedin.com/in/kishorereddy945
                </a>
            </p>
            <p style="font-size:0.85em;">
                <a href="https://github.com/dkreddy2025" style="color:#22C55E; display:flex; align-items:center; gap:8px; text-decoration:none;" target="_blank">
                    <img src="https://cdn.jsdelivr.net/npm/simple-icons@v9/icons/github.svg" width="16" style="filter:invert(1) sepia(1) saturate(5) hue-rotate(90deg);"/>
                    github.com/dkreddy2025
                </a>
            </p>
        </div>
    </div>
    <div style="text-align:center; margin-top:30px; padding-top:15px; border-top:1px solid #334155;">
        <p style="color:#94A3B8; font-size:0.8em;">© 2026 SportsIQ AI • </p>
        <p style="color:#94A3B8; font-size:0.8em;">Designed & Developed by Kishore Reddy Dudipala</p>
    </div>
</div>
""", unsafe_allow_html=True)